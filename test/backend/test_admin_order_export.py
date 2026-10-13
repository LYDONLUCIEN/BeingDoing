"""
Admin 营收汇总 + 订单导出（xlsx）测试（2026-10-13）

覆盖：
- 路由层（TestClient + mock 服务层，仿 test_refund_api.py）：
  非超管 403；export 返回 xlsx 且可被 openpyxl 读回、筛选透传、写审计、
  超行数上限 400；revenue-summary 200
- 服务层（in-memory SQLite，仿 test_payment_service.py）：
  admin_export_orders 筛选/无分页/行数上限；
  admin_revenue_summary 净实收口径（Σ实付 − Σ已退款、排除 pending/closed）、
  今日/本月 +08:00 日界

审计写入 mock 到 tmp_path，不污染真实 data/。
"""

import io
from datetime import datetime, timedelta, timezone
from unittest.mock import AsyncMock, patch

import pytest
from app.api.v1.auth import get_current_user
from app.main import app
from app.models.database import Base
from app.models.payment import PaymentOrder
from app.models.user import User
from app.services import order_export_service as oes_mod
from app.services import payment_service as ps_mod
from app.services.payment_service import PaymentService
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

_USER = {"user_id": "u1", "email": "alice@test.com"}
_ADMIN = {"user_id": "admin-1", "email": "admin@test.com"}

SH_TZ = timezone(timedelta(hours=8))

# ─── 路由层 ──────────────────────────────────────────────────


def _client(user: dict) -> TestClient:
    app.dependency_overrides[get_current_user] = lambda: user
    return TestClient(app)


@pytest.fixture(autouse=True)
def _clear_overrides():
    yield
    app.dependency_overrides.clear()


def _as_admin():
    return patch("app.api.v1.admin_payment.is_super_admin_user", return_value=True)


_SAMPLE_ITEMS = [
    {
        "order_no": "X202610130001",
        "user_email": "alice@test.com",
        "product_type": "quarterly_package",
        "quantity": 1,
        "amount_original": 19900,
        "amount_discount": 1000,
        "amount_paid": 18900,
        "amount_refunded": 0,
        "coupon_code": "ABC123",
        "channel": "alipay",
        "status": "granted",
        "delivered_code": "CODE-1",
        "channel_transaction_id": "T1",
        "created_at": "2026-10-13T02:30:00",
        "paid_at": "2026-10-13T02:31:00",
        "closed_at": None,
        "refunded_at": None,
    },
    {
        "order_no": "X202610130002",
        "user_email": "bob@test.com",
        "product_type": "consultation",
        "quantity": 1,
        "amount_original": 39900,
        "amount_discount": 0,
        "amount_paid": 39900,
        "amount_refunded": 39900,
        "coupon_code": None,
        "channel": "alipay",
        "status": "refunded",
        "delivered_code": None,
        "channel_transaction_id": "T2",
        "created_at": "2026-10-12T03:00:00",
        "paid_at": "2026-10-12T03:01:00",
        "closed_at": None,
        "refunded_at": "2026-10-13T01:00:00",
    },
]


def test_export_and_summary_require_super_admin():
    """非超管访问导出 / 营收汇总 → 403"""
    client = _client(_USER)
    assert client.get("/api/v1/admin/payment/orders/export").status_code == 403
    assert client.get("/api/v1/admin/payment/revenue-summary").status_code == 403


def test_export_orders_xlsx_ok():
    """导出路由：200 + xlsx content-type + 筛选透传 + 审计写入（mock 到 tmp_path）"""
    audit_calls = []

    def _fake_audit(actor_email, filters, row_count):
        audit_calls.append((actor_email, filters, row_count))

    client = _client(_ADMIN)
    with _as_admin(), patch.object(
        PaymentService,
        "admin_export_orders",
        new_callable=AsyncMock,
        return_value=list(_SAMPLE_ITEMS),
    ) as mock_query, patch(
        "app.api.v1.admin_payment.append_order_export_audit", side_effect=_fake_audit
    ):
        res = client.get("/api/v1/admin/payment/orders/export?status=granted")

    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "orders_" in res.headers["content-disposition"]
    assert ".xlsx" in res.headers["content-disposition"]

    # 筛选透传（未传的 channel 为 None）
    mock_query.assert_awaited_once_with(status="granted", channel=None)

    # 审计：actor 邮箱 + 筛选 + 行数
    assert audit_calls == [("admin@test.com", {"status": "granted", "channel": None}, 2)]

    # xlsx 可读回：表头 + 2 行，金额元、中文标签、时间 +08:00
    ws = load_workbook(io.BytesIO(res.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "订单号"
    assert len(rows) == 3
    assert rows[1][1] == "alice@test.com"
    assert rows[1][2] == "季度套餐"
    assert rows[1][6] == 189.0  # 实付(元)
    assert rows[1][9] == "已发放"
    assert rows[1][13] == "2026-10-13 10:30"  # 02:30 UTC → 10:30 +08:00
    assert rows[2][9] == "已退款"


def test_export_orders_over_limit_400():
    """超出行数上限 → 400（服务层 ValueError 映射）"""
    client = _client(_ADMIN)
    with _as_admin(), patch.object(
        PaymentService,
        "admin_export_orders",
        new_callable=AsyncMock,
        side_effect=ValueError("导出超过 20000 行上限，请先加筛选条件分批导出"),
    ):
        res = client.get("/api/v1/admin/payment/orders/export")
    assert res.status_code == 400
    assert "上限" in res.json()["detail"]


def test_revenue_summary_ok():
    """营收汇总路由：200 + 契约字段"""
    summary = {
        "total_paid_cents": 100,
        "total_refunded_cents": 20,
        "net_revenue_cents": 80,
        "paid_order_count": 2,
        "today_net_cents": 50,
        "month_net_cents": 80,
        "generated_at": "2026-10-13T00:00:00+00:00",
    }
    client = _client(_ADMIN)
    with _as_admin(), patch.object(
        PaymentService,
        "admin_revenue_summary",
        new_callable=AsyncMock,
        return_value=summary,
    ):
        res = client.get("/api/v1/admin/payment/revenue-summary")
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["net_revenue_cents"] == 80
    assert data["paid_order_count"] == 2


# ─── 服务层（in-memory SQLite）────────────────────────────────

_test_engine = create_async_engine("sqlite+aiosqlite:///:memory:", future=True)
_TestSessionLocal = async_sessionmaker(_test_engine, expire_on_commit=False)


@pytest.fixture(autouse=True)
async def _setup_db(monkeypatch, tmp_path):
    """建表 + 插测试用户；AsyncSessionLocal 指向内存库；审计路径指向 tmp_path"""
    async with _test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    monkeypatch.setattr(ps_mod, "AsyncSessionLocal", _TestSessionLocal)
    monkeypatch.setattr(
        oes_mod, "_audit_path", lambda: tmp_path / "admin_export_audit.jsonl"
    )

    now = datetime.now(timezone.utc)
    async with _TestSessionLocal() as db:
        db.add(
            User(
                id="u1",
                email="alice@test.com",
                username="alice",
                password_hash="x",
                is_active=True,
                created_at=now,
            )
        )
        await db.commit()

    yield

    async with _test_engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)


def _sh_naive_utc(dt_sh: datetime) -> datetime:
    """+08:00 时刻 → DB 存储的 naive UTC"""
    return dt_sh.astimezone(timezone.utc).replace(tzinfo=None)


async def _insert_order(
    order_no: str,
    *,
    status: str,
    amount_paid: int,
    amount_refunded: int = 0,
    channel: str = "alipay",
    paid_at: datetime | None = None,
    created_at: datetime | None = None,
) -> None:
    async with _TestSessionLocal() as db:
        db.add(
            PaymentOrder(
                order_no=order_no,
                user_id="u1",
                product_type="quarterly_package",
                quantity=1,
                amount_original=amount_paid,
                amount_discount=0,
                amount_paid=amount_paid,
                amount_refunded=amount_refunded,
                channel=channel,
                status=status,
                paid_at=paid_at,
                created_at=created_at or datetime.now(timezone.utc).replace(tzinfo=None),
            )
        )
        await db.commit()


async def test_admin_export_orders_filters_and_fields():
    """导出查询：status/channel 筛选生效、无分页、字段对齐（含 user_email，无 code_refundable）"""
    await _insert_order("E1", status="granted", amount_paid=100, channel="alipay")
    await _insert_order("E2", status="pending", amount_paid=200, channel="alipay")
    await _insert_order("E3", status="granted", amount_paid=300, channel="wechat")

    all_items = await PaymentService.admin_export_orders()
    assert len(all_items) == 3
    item = next(i for i in all_items if i["order_no"] == "E1")
    assert item["user_email"] == "alice@test.com"
    assert "code_refundable" not in item
    assert item["amount_paid"] == 100

    granted = await PaymentService.admin_export_orders(status="granted")
    assert {i["order_no"] for i in granted} == {"E1", "E3"}

    wechat = await PaymentService.admin_export_orders(channel="wechat")
    assert [i["order_no"] for i in wechat] == ["E3"]


async def test_admin_export_orders_row_cap(monkeypatch):
    """导出行数上限：超出抛 ValueError（路由转 400）"""
    monkeypatch.setattr(ps_mod, "_ADMIN_EXPORT_MAX_ROWS", 2)
    for i in range(3):
        await _insert_order(f"C{i}", status="granted", amount_paid=100)
    with pytest.raises(ValueError, match="上限"):
        await PaymentService.admin_export_orders()


async def test_admin_revenue_summary_net_amounts():
    """净实收口径：Σ实付 − Σ已退款；pending/closed 不计；refunded 净额为 0"""
    now_naive = datetime.now(timezone.utc).replace(tzinfo=None)
    await _insert_order("R1", status="granted", amount_paid=10000, paid_at=now_naive)
    await _insert_order(
        "R2", status="refunded", amount_paid=5000, amount_refunded=5000, paid_at=now_naive
    )
    await _insert_order(
        "R3",
        status="partially_refunded",
        amount_paid=8000,
        amount_refunded=3000,
        paid_at=now_naive,
    )
    await _insert_order("R4", status="pending", amount_paid=19900)  # 不计
    await _insert_order("R5", status="closed", amount_paid=19900)  # 不计

    s = await PaymentService.admin_revenue_summary()
    assert s["total_paid_cents"] == 23000
    assert s["total_refunded_cents"] == 8000
    assert s["net_revenue_cents"] == 15000
    assert s["paid_order_count"] == 3
    # 全部 paid_at 在当下 → 今日/本月同累计
    assert s["today_net_cents"] == 15000
    assert s["month_net_cents"] == 15000


async def test_admin_revenue_summary_shanghai_day_boundary():
    """今日/本月按 +08:00 日界归属 paid_at"""
    now_sh = datetime.now(timezone.utc).astimezone(SH_TZ)
    today_start_sh = now_sh.replace(hour=0, minute=0, second=0, microsecond=0)
    month_start_sh = today_start_sh.replace(day=1)

    # 今日 00:01（+08:00）→ 计今日
    await _insert_order(
        "D1",
        status="granted",
        amount_paid=1000,
        paid_at=_sh_naive_utc(today_start_sh + timedelta(minutes=1)),
    )
    # 昨日 23:59（+08:00）→ 不计今日；是否计本月取决于今天是否 1 号
    await _insert_order(
        "D2",
        status="granted",
        amount_paid=2000,
        paid_at=_sh_naive_utc(today_start_sh - timedelta(minutes=1)),
    )
    # 上月末 23:59（+08:00）→ 不计本月（也不计今日）
    await _insert_order(
        "D3",
        status="granted",
        amount_paid=4000,
        paid_at=_sh_naive_utc(month_start_sh - timedelta(minutes=1)),
    )

    s = await PaymentService.admin_revenue_summary()
    assert s["today_net_cents"] == 1000
    d2_in_month = today_start_sh.day != 1
    assert s["month_net_cents"] == 1000 + (2000 if d2_in_month else 0)
    assert s["net_revenue_cents"] == 7000
