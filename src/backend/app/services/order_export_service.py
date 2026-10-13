"""
Admin 订单明细导出（xlsx）与导出审计

- build_orders_xlsx：PaymentService.admin_export_orders 返回的 dict 列表 → xlsx 字节
  （openpyxl 内存构建；金额分→元，状态/渠道/商品中文化，时间转 +08:00 本地）
- append_order_export_audit：导出行为审计（含用户邮箱属个人敏感信息，需留痕），
  append-only JSONL 写 data/admin_export_audit.jsonl，写失败仅记日志不阻断
  （模式仿 account_deletion_service.append_account_deletion_audit，独立文件不混注销语义）
"""

from __future__ import annotations

import io
import json
import logging
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook

from app.utils.data_paths import get_project_data_dir

logger = logging.getLogger(__name__)

# Asia/Shanghai 固定 UTC+8（无夏令时），与 analytics_funnel_service 口径一致
SH_TZ = timezone(timedelta(hours=8))

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADERS = [
    "订单号",
    "用户邮箱",
    "商品",
    "数量",
    "原价(元)",
    "抵扣(元)",
    "实付(元)",
    "已退款(元)",
    "渠道",
    "状态",
    "使用券码",
    "交付激活码",
    "渠道交易号",
    "创建时间",
    "支付时间",
    "关闭时间",
    "退款时间",
]

_STATUS_LABEL = {
    "pending": "待支付",
    "paid": "已支付",
    "granted": "已发放",
    "closed": "已关闭",
    "cancelled": "已取消",
    "refunding": "退款中",
    "partially_refunded": "部分退款",
    "refunded": "已退款",
}

_CHANNEL_LABEL = {"alipay": "支付宝", "wechat": "微信"}

_PRODUCT_LABEL = {
    "quarterly_package": "季度套餐",
    "annual_package": "年度套餐",
    "renewal": "延期激活",
    "consultation": "报告解读咨询",
    # 旧 SKU / 会员预留（历史订单兼容）
    "activation_code": "全程激活码",
    "membership_monthly": "会员（包月）",
    "membership_lifetime": "会员（永久）",
}

# 金额列（HEADERS 下标，0 起）：设置 '0.00' 数字格式
_AMOUNT_COLS = (4, 5, 6, 7)


def _fen_to_yuan(value: Optional[int]) -> float:
    return round((value or 0) / 100, 2)


def _fmt_time(iso_str: Optional[str]) -> str:
    """ISO 时间（DB naive 按 UTC 解释）→ +08:00 本地 'YYYY-MM-DD HH:MM'"""
    if not iso_str:
        return ""
    try:
        dt = datetime.fromisoformat(iso_str)
    except ValueError:
        return iso_str
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(SH_TZ).strftime("%Y-%m-%d %H:%M")


def build_orders_xlsx(items: List[Dict[str, Any]]) -> bytes:
    """订单 dict 列表 → xlsx 字节（一行一订单，不展开 order_lines）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "订单明细"
    ws.append(HEADERS)
    for it in items:
        ws.append(
            [
                it.get("order_no") or "",
                it.get("user_email") or "",
                _PRODUCT_LABEL.get(it.get("product_type") or "", it.get("product_type") or ""),
                it.get("quantity") or 1,
                _fen_to_yuan(it.get("amount_original")),
                _fen_to_yuan(it.get("amount_discount")),
                _fen_to_yuan(it.get("amount_paid")),
                _fen_to_yuan(it.get("amount_refunded")),
                _CHANNEL_LABEL.get(it.get("channel") or "", it.get("channel") or ""),
                _STATUS_LABEL.get(it.get("status") or "", it.get("status") or ""),
                it.get("coupon_code") or "",
                it.get("delivered_code") or "",
                it.get("channel_transaction_id") or "",
                _fmt_time(it.get("created_at")),
                _fmt_time(it.get("paid_at")),
                _fmt_time(it.get("closed_at")),
                _fmt_time(it.get("refunded_at")),
            ]
        )
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for col in _AMOUNT_COLS:
            row[col].number_format = "0.00"
    # 适度列宽（中文按 2 字符宽估算）
    for idx, header in enumerate(HEADERS, start=1):
        max_len = len(header) * 2
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=idx, max_col=idx):
            cell = row[0]
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)) * 2, 60))
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = max_len + 2
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_export_filename(now: Optional[datetime] = None) -> str:
    """导出文件名（+08:00 本地时间戳）"""
    dt = (now or datetime.now(timezone.utc)).astimezone(SH_TZ)
    return f"orders_{dt.strftime('%Y%m%d_%H%M%S')}.xlsx"


def _audit_path() -> Path:
    return get_project_data_dir() / "admin_export_audit.jsonl"


def append_order_export_audit(
    actor_email: Optional[str],
    filters: Dict[str, Any],
    row_count: int,
) -> None:
    """追加一条订单导出审计日志（JSONL，写失败只记日志不阻断）"""
    entry = {
        "event": "order_export",
        "actor_email": actor_email,
        "filters": filters,
        "row_count": row_count,
        "timestamp": datetime.now(timezone.utc).isoformat(),
    }
    p = _audit_path()
    try:
        p.parent.mkdir(parents=True, exist_ok=True)
        with p.open("a", encoding="utf-8") as f:
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")
    except OSError:
        logger.exception("写入订单导出审计日志失败: actor=%s", actor_email)
